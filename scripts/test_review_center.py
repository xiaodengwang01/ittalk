"""Private review artifacts with simulated contacts only. No network or real mail."""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from scripts import review_center as center
from scripts import test_daily_workflow as daily_tests
from scripts import test_mail_send as send_tests
from docx import Document


class ReviewTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.shared = tempfile.TemporaryDirectory()
        cls.fixture = Path(cls.shared.name) / 'fixture.xlsx'
        daily_tests.fixture(cls.fixture)

    @classmethod
    def tearDownClass(cls): cls.shared.cleanup()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.book = self.root / 'business/book.xlsx'
        self.book.parent.mkdir()
        self.book.write_bytes(self.fixture.read_bytes())
        self.before = self.book.read_bytes()
        self.state = center.mail.State(self.root / 'private/state.sqlite3')
        center.initialize(self.state)
        self.incoming()

    def tearDown(self):
        self.assertEqual(self.book.read_bytes(), self.before)
        self.state.close(); self.temp.cleanup()

    def incoming(self, identifier='<sim-in@example.invalid>', text='How much is the trial?', date='2026-10-21T10:00:00+08:00'):
        center.ledger.incoming(self.state, {'account': 'team@example.invalid', 'message_id': identifier,
            'creator_id': 'SIM-000001', 'source': 'INBOX:UID='+identifier, 'received_at': date, 'raw_text': text}, 'SIM-000001')

    def report(self, body='Hi Sim,\n\nComplete draft\n\nBest,\nSara', image=False):
        file = self.root / 'source.docx'
        doc = Document()
        doc.add_paragraph('模拟带教汇报', 'Title')
        doc.add_heading('模拟达人合作情况', 1)
        doc.add_paragraph('邮箱 sim1@example.invalid')
        doc.add_paragraph('请带教决定是否接受条件。')
        doc.add_paragraph('收件人：sim1@example.invalid')
        doc.add_paragraph('主题：Re: simulated collaboration')
        doc.add_paragraph(body)
        if image:
            import base64
            png = self.root / 'image.png'
            png.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j6ioAAAAASUVORK5CYII='))
            doc.add_picture(str(png))
        doc.add_paragraph('审核状态：待确认，未发送。')
        doc.save(file)
        return file

    def build(self):
        return center.build(self.state, '2026-10-22', self.book)

    def test_empty_daily_view_does_not_invent_new_creators_or_send(self):
        result = self.build()
        self.assertEqual(result['counts'], {'added': 0, 'mentor': 0, 'drafts': 0})
        self.assertFalse(result['sent'])
        self.assertIn('尚无已登记的新增达人', Path(result['page']).read_text())
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM send_attempts').fetchone()[0], 0)

    def test_document_adoption_preserves_manual_text_full_reply_and_images(self):
        source = self.report('Hi Sim,\n\nManually revised complete reply\n\nBest,\nSara', image=True)
        raw = source.read_bytes()
        result = center.adopt_document(self.state, source, self.book)
        self.assertEqual(result['imported'], ['SIM-000001'])
        report = self.build()
        with ZipFile(report['document']) as z:
            self.assertEqual(len([n for n in z.namelist() if n.startswith('word/media/')]), 1)
        text = '\n'.join(p.text for p in Document(report['document']).paragraphs)
        self.assertIn('Manually revised complete reply', text)
        self.assertEqual(source.read_bytes(), raw)
        packet = center.packets(self.state)['SIM-000001']
        images = [i for b in packet['draft']['blocks'] for i in b.get('images', [])]
        self.assertTrue(all(Path(i['path']).is_relative_to(self.state.path.parent) for i in images))

    def test_same_input_reuses_document_and_keeps_send_approval_absent(self):
        center.adopt_document(self.state, self.report(), self.book)
        one = self.build()
        body = Path(one['document']).read_bytes()
        two = self.build()
        self.assertEqual(one['model_hash'], two['model_hash'])
        self.assertEqual(body, Path(two['document']).read_bytes())
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM send_reviews').fetchone()[0], 0)

    def test_pool_only_changes_refresh_page_without_reexporting_mentor_document(self):
        import candidate_pool
        center.adopt_document(self.state,self.report(),self.book)
        first=self.build();document=Path(first['document']).read_bytes()
        candidate_pool.ingest(self.state,{'creators':[{'账号链接':'https://tiktok.com/@sim_pool'}]},self.book)
        second=self.build()
        self.assertNotEqual(first['model_hash'],second['model_hash'])
        self.assertEqual(document,Path(second['document']).read_bytes())
        self.assertIn('@sim_pool',Path(second['page']).read_text())

    def test_failed_research_routes_are_visible_without_freezing_the_day(self):
        center.daily.put(self.state, 'day:2026-10-22', {'creators': [], 'prepared': False})
        center.daily.put(self.state, 'pool-routes:2026-10-22', {
            'routes': [], 'limitations': ['模拟来源读取超时，未取得当天视频证据。']})
        result = self.build()
        text = '\n'.join(p.text for p in Document(result['document']).paragraphs)
        for content in (text, Path(result['page']).read_text()):
            self.assertIn('模拟来源读取超时', content)
            self.assertIn('仍缺 5 位', content)
            self.assertIn('搜集尚未完成', content)
        self.assertFalse(center.daily.get(self.state, 'day:2026-10-22')['prepared'])
        self.assertEqual(result['counts']['added'], 0)

    def test_latest_pending_summary_takes_precedence_over_old_workbook_text(self):
        rows, appendix = center.source_rows(self.book)
        rows['SIM-000001']['沟通摘要'] = '旧表格摘要'
        rows['SIM-000001']['建议回复策略'] = '旧策略'
        center.daily.put(self.state, 'reply:sim-current', {'message': {
            'creator_id': 'SIM-000001', 'message_id': '<sim-in@example.invalid>',
            'received_at': '2026-10-21T10:00:00+08:00', 'summary': '最新反报价摘要',
            'reason': '最新推进方向', 'decisions': ['最新待决事项']}})
        with patch.object(center, 'source_rows', return_value=(rows, appendix)):
            card = center.model(self.state, '2026-10-22', self.book)['cards'][0]
        self.assertEqual(card['summary'], '最新反报价摘要')
        self.assertEqual(card['strategy'], '最新推进方向')
        self.assertEqual(card['decisions'], ['最新待决事项'])

    def test_generated_doc_manual_changes_are_adopted_before_refresh(self):
        center.adopt_document(self.state, self.report(), self.book)
        one = self.build()
        doc = Document(one['document'])
        paragraph = next(p for p in doc.paragraphs if 'Complete draft' in p.text)
        paragraph.text = 'Hi Sim,\n\nHuman changed after export\n\nBest,\nSara'
        doc.save(one['document'])
        two = self.build()
        text = '\n'.join(p.text for p in Document(two['document']).paragraphs)
        self.assertIn('Human changed after export', text)
        self.assertIn('Human changed after export', Path(two['page']).read_text())

    def test_cc_bcc_are_in_both_document_and_page_and_multiple_primary_recipients_rejected(self):
        source = self.report()
        doc = Document(source)
        subject = next(p for p in doc.paragraphs if p.text.startswith('主题：'))
        for value in ('抄送：copy@example.invalid', '密送：private@example.invalid'):
            paragraph = doc.add_paragraph(value)
            subject._p.addprevious(paragraph._p)
        doc.save(source)
        center.adopt_document(self.state, source, self.book)
        result = self.build()
        content = '\n'.join(p.text for p in Document(result['document']).paragraphs)
        for address in ('copy@example.invalid', 'private@example.invalid'):
            self.assertIn(address, content)
            self.assertIn(address, Path(result['page']).read_text())
        recipient = next(p for p in doc.paragraphs if p.text.startswith('收件人：'))
        recipient.text += ', unknown@example.invalid'
        doc.save(source)
        with self.assertRaises(center.mail.DraftError): center.adopt_document(self.state, source, self.book)

    def test_edit_during_export_is_preserved(self):
        import review_docx
        center.adopt_document(self.state, self.report(), self.book)
        first = self.build()
        self.incoming('<sim-new@example.invalid>', 'New terms', '2026-10-22T11:00:00+08:00')
        writer = review_docx.write_document
        def competing_edit(snapshot, temporary):
            writer(snapshot, temporary)
            Path(first['document']).write_bytes(b'human edit during export')
        with patch.object(review_docx, 'write_document', side_effect=competing_edit):
            with self.assertRaises(center.mail.DraftError): self.build()
        self.assertEqual(Path(first['document']).read_bytes(), b'human edit during export')

    def test_new_message_invalidates_old_reply_and_format_edit_does_not_reapprove_it(self):
        source = self.report('OLD DRAFT MUST NOT BE SHOWN')
        center.adopt_document(self.state, source, self.book)
        self.incoming('<sim-new@example.invalid>', 'Changed price', '2026-10-22T11:00:00+08:00')
        doc = Document(source);doc.paragraphs[0].text = 'New formatting';doc.save(source)
        result = self.build()
        self.assertEqual(result['counts']['drafts'], 0)
        self.assertNotIn('OLD DRAFT MUST NOT BE SHOWN', Path(result['page']).read_text())
        self.assertTrue(result['warnings'])

    def test_reply_to_latest_message_removes_waiting_card(self):
        center.adopt_document(self.state, self.report(), self.book)
        center.ledger.put(self.state, {'account': 'team@example.invalid', 'message_id': '<sim-out@example.invalid>',
            'direction': 'outgoing', 'source': 'Sent:1', 'creator_id': 'SIM-000001',
            'sent_at': '2026-10-21T11:00:00+08:00', 'in_reply_to': '<sim-in@example.invalid>'})
        self.assertEqual(center.model(self.state, '2026-10-22', self.book)['cards'], [])

    def test_changed_or_missing_image_blocks_complete_preview(self):
        center.adopt_document(self.state, self.report(image=True), self.book)
        packet = center.packets(self.state)['SIM-000001']
        image = next(i for b in packet['draft']['blocks'] for i in b.get('images', []))
        Path(image['path']).unlink()
        result = self.build()
        self.assertEqual(result['counts']['drafts'], 0)
        self.assertTrue(result['warnings'])

    def test_ambiguous_report_or_placeholder_does_not_replace_existing_packet(self):
        source = self.report()
        center.adopt_document(self.state, source, self.book)
        previous = center.packets(self.state)
        for body in ('Hello {{missing}}', ''):
            source = self.report(body)
            with self.assertRaises(center.mail.DraftError): center.adopt_document(self.state, source, self.book)
            self.assertEqual(previous, center.packets(self.state))

    def test_foreign_existing_document_is_not_overwritten(self):
        target = self.book.parent / '参考资料/2026-10-22_KOL跟进.docx'
        target.parent.mkdir()
        target.write_bytes(b'untouched human document')
        with self.assertRaises(center.mail.DraftError): self.build()
        self.assertEqual(target.read_bytes(), b'untouched human document')

    def test_hostile_email_is_escaped_and_page_has_no_send_or_network_action(self):
        center.adopt_document(self.state, self.report('<script>fetch("https://example.invalid")</script>'), self.book)
        result = self.build()
        page = Path(result['page']).read_text()
        self.assertIn('&lt;script&gt;fetch', page)
        self.assertNotIn('<script>fetch', page)
        self.assertIn("default-src &#x27;none&#x27;", page)
        self.assertNotIn('/api/send', page)
        self.assertNotIn('method="post"', page)

    def test_artifact_failure_does_not_disguise_completed_business_operation(self):
        with patch.object(center, 'build', side_effect=OSError('simulated')):
            value = center.refresh_after_operation(self.state, {'recorded': ['SIM-000001']}, self.book)
        self.assertEqual(value['status'], 'needs_refresh')


class SentReviewTests(unittest.TestCase):
    setUp = send_tests.SendTests.setUp
    tearDown = send_tests.SendTests.tearDown
    review = send_tests.SendTests.review
    approve = send_tests.SendTests.approve
    edit = send_tests.SendTests.edit

    def test_final_mime_preserves_attachments_and_sandboxes_html(self):
        center.initialize(self.state)
        self.edit(html='<p>Full HTML version</p><script>alert(1)</script>', attachment=b'complete attachment')
        self.review()
        rows = {f'SIM-{i:03d}': {'达人ID': f'SIM-{i:03d}', 'Email': f'sim{i}@example.invalid'} for i in (1, 2)}
        book = self.folder / 'book.xlsx';book.write_bytes(b'simulated')
        with patch.object(center, 'source_rows', return_value=(rows, {})):
            snapshot = center.model(self.state, '2026-10-22', book)
        from review_page import render_page
        page = render_page(snapshot, self.folder / 'report.docx')
        self.assertEqual(len(snapshot['mails']), 2)
        self.assertIn('sample.bin', page)
        self.assertIn('sandbox=""', page)
        self.assertNotIn('<script>alert(1)</script>', page)
        self.assertIn(center.sender.base64.b64encode(b'complete attachment').decode(), page)

    def test_sent_review_and_its_old_waiting_copy_never_reappear_as_pending(self):
        center.initialize(self.state)
        review = self.review()
        old = deepcopy(review); old['id'] = 'sim-old-waiting'
        center.sender.put_review(self.state, old)
        self.approve(review)
        transport = send_tests.FakeTransport(self.mailbox)
        center.sender.deliver(self.state, review['id'], self.mailbox, lambda _: transport, pause=lambda _: None)
        rows = {f'SIM-{i:03d}': {'达人ID': f'SIM-{i:03d}', 'Email': f'sim{i}@example.invalid'} for i in (1, 2)}
        with patch.object(center, 'source_rows', return_value=(rows, {})):
            book = self.folder / 'book.xlsx';book.write_bytes(b'simulated')
            value = center.model(self.state, '2026-10-22', book)
        self.assertEqual(value['mails'], [])
        self.assertEqual(value['verified_total'], 2)


if __name__ == '__main__': unittest.main()
