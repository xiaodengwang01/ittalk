"""Compact workbook workflow regression tests, using only synthetic contacts."""
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from scripts import review_center as c
from scripts import mail_postprocess as post
from scripts import candidate_pool as pool
from scripts import test_daily_workflow as support
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

w=c.daily.workbook


class CompactTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.shared=tempfile.TemporaryDirectory()
        root=Path(cls.shared.name)
        cls.fixture=root/'book.xlsx'
        rows=[]
        for i in (1,2):
            rows.append({'达人ID':f'SIM-{i:06d}','达人昵称':f'模拟达人 {i}',
                         'Email':f'sim{i}@example.invalid','当前建联状况':'收到新回信' if i==2 else '待首联确认',
                         '账号链接':f'https://www.tiktok.com/@sim_{i}','首联内容分类':'宠物剧情类',
                         '往来邮件梗概':'2026-10-21｜对方：询问试用周期。' if i==2 else '',
                         '内容能力(1-5)':4,'宠物人群匹配度(1-5)':5,'播放能力(1-5)':4,'联系方式完整度(1-5)':5,
                         '_scoreFormula':True})
        for h in ('内容能力(1-5)','宠物人群匹配度(1-5)','播放能力(1-5)','联系方式完整度(1-5)'):
            rows[0][h]=None
        spec={'headers':w.COMPACT_HEADERS,'sheets':[{'name':name,'rows':[row]} for name,row in zip(w.SHEETS,rows)]}
        source=root/'spec.json';source.write_text(json.dumps(spec))
        node=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'
        subprocess.run([str(node),str(Path(w.__file__).with_name('followup_workbook.mjs')),str(source),str(cls.fixture)],check=True,capture_output=True,timeout=120)
        w.compact_controls(cls.fixture)

    @classmethod
    def tearDownClass(cls):cls.shared.cleanup()

    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.book=self.root/'book.xlsx';self.book.write_bytes(self.fixture.read_bytes())
        self.state=c.mail.State(self.root/'state.sqlite3');c.initialize(self.state)
        self.incoming={'account':'team@example.invalid','message_id':'<in@example.invalid>','creator_id':'SIM-000002',
                       'source':'INBOX:UID=2','received_at':'2026-10-21T10:00:00+08:00','raw_text':'Can we test for a week?'}
        c.ledger.incoming(self.state,self.incoming,'SIM-000002')

    def tearDown(self):self.state.close();self.temp.cleanup()

    def report(self, *, revisions=False, comment=False):
        path=self.root/'mentor.docx';doc=Document();doc.add_heading('模拟达人 2',1)
        doc.add_paragraph('达人ID：SIM-000002')
        doc.add_heading('对方来信原文',2);doc.add_paragraph('主题：Incoming subject')
        doc.add_paragraph('收件人：team@example.invalid');doc.add_paragraph(self.incoming['raw_text'])
        doc.add_heading('来信要点',2);doc.add_paragraph('询问能否试用一周。')
        doc.add_heading('我方回复要点及应对策略',2);doc.add_paragraph('提议一周试用，具体安排待确认。')
        doc.add_heading('我方完整回复正文',2);doc.add_paragraph('收件人：sim2@example.invalid')
        doc.add_paragraph('主题：Re: simulated trial')
        p=doc.add_paragraph('We suggest a one-week trial.')
        if revisions:
            p.text=''
            for tag,text in [('del','We promise a paid contract.'),('ins','We suggest a one-week trial.')]:
                node=OxmlElement('w:'+tag);node.set(qn('w:id'),'1' if tag=='del' else '2')
                run=OxmlElement('w:r');t=OxmlElement('w:delText' if tag=='del' else 'w:t');t.text=text
                run.append(t);node.append(run);p._p.append(node)
        if comment:
            doc.add_comment(p.add_run(' Please confirm.'),text='仅提议试用，不承诺签约。',author='Mentor')
        doc.add_paragraph('审核状态：待确认，未发送。');doc.save(path)
        return path

    def updates(self,source,comments=False):
        return {'source_hash':c.mail.digest(source.read_bytes()),'book_hash':c.mail.digest(self.book.read_bytes()),
                'creators':{'SIM-000002':{'comments_reviewed':comments,'fields':{
                    '待发送回复梗概':'未发送｜提议一周试用，具体安排待确认。','当前建联状况':'回复稿待确认',
                    'WhatsApp联系方式':'不使用，继续邮件沟通','收件信息':'姓名：未提供\n地址：未提供'}}}}

    def event(self,body='We suggest a one-week trial.'):
        return {'account':'team@example.invalid','creator_id':'SIM-000002','message_id':'<out@example.invalid>',
                'sent_at':'2026-10-22T11:00:00+08:00','to':'sim2@example.invalid','subject':'Re: simulated trial',
                'in_reply_to':'<in@example.invalid>','raw_text':body}

    def test_document_update_and_verified_send_use_secondary_and_never_store_body(self):
        source=self.report();before=source.read_bytes()
        result=c.adopt_document(self.state,source,self.book,self.updates(source))
        self.assertFalse(result['sent']);self.assertEqual(source.read_bytes(),before)
        sheet,row=w.locate(self.book,'SIM-000002')
        self.assertEqual(sheet['name'],w.SHEETS[1]);self.assertTrue(row['待发送回复梗概'].startswith('未发送｜'))
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM send_attempts').fetchone()[0],0)
        entry={'creator_id':'SIM-000002','batch_id':'sim-batch'};review={'account':'team@example.invalid'}
        post.write_workbook(self.state,c.sender,review,entry,self.event(),self.book)
        _,row=w.locate(self.book,'SIM-000002')
        self.assertFalse(row['待发送回复梗概'])
        self.assertEqual(row['往来邮件梗概'].count('2026-10-22｜我方：提议一周试用'),1)
        raw=self.book.read_bytes();post.write_workbook(self.state,c.sender,review,entry,self.event(),self.book)
        self.assertEqual(self.book.read_bytes(),raw)
        with ZipFile(self.book) as z:
            for n in z.namelist():self.assertNotIn(b'We suggest a one-week trial.',z.read(n))
        self.assertFalse((self.root/'backups').exists())
        _,main=w.locate(self.book,'SIM-000001');self.assertFalse(main.get('往来邮件梗概'))

    def test_changed_sent_body_does_not_reuse_or_clear_older_pending_synopsis(self):
        source=self.report();c.adopt_document(self.state,source,self.book,self.updates(source))
        post.write_workbook(self.state,c.sender,{'account':'team@example.invalid'},
            {'creator_id':'SIM-000002','batch_id':'sim'},self.event('Different actual sent terms.'),self.book)
        _,row=w.locate(self.book,'SIM-000002')
        self.assertTrue(row['待发送回复梗概'].startswith('未发送｜'))
        self.assertIn('内容梗概待整理',row['往来邮件梗概'])

    def test_revisions_and_comments_are_read_without_altering_original(self):
        source=self.report(revisions=True,comment=True);before=source.read_bytes()
        result=c.adopt_document(self.state,source,self.book)
        self.assertIn('SIM-000002',result['comments'])
        packet=c.packets(self.state)['SIM-000002'];body=str(packet['draft'])
        self.assertIn('We suggest a one-week trial.',body);self.assertNotIn('We promise',body)
        self.assertEqual(packet['draft']['subject'],'Re: simulated trial')
        with self.assertRaises(c.mail.DraftError):c.adopt_document(self.state,source,self.book,self.updates(source))
        c.adopt_document(self.state,source,self.book,self.updates(source,True))
        self.assertEqual(source.read_bytes(),before)
        self.assertFalse(c.packets(self.state)['SIM-000002']['unresolved_comments'])

    def test_changed_source_or_workbook_rejects_stale_updates(self):
        source=self.report();updates=self.updates(source)
        w.update_fields(self.book,'SIM-000002',{'沟通摘要':'人工最新内容'},expected_hash=c.mail.digest(self.book.read_bytes()))
        with self.assertRaises(c.mail.DraftError):c.adopt_document(self.state,source,self.book,updates)
        self.assertEqual(w.locate(self.book,'SIM-000002')[1]['沟通摘要'],'人工最新内容')

    def test_new_incoming_summary_is_appended_once_by_message_id(self):
        new={**self.incoming,'message_id':'<new@example.invalid>','received_at':'2026-10-22T13:00:00+08:00','raw_text':'Please clarify return shipping.'}
        c.ledger.incoming(self.state,new,'SIM-000002')
        source=self.report()
        for _ in range(2):
            updates=self.updates(source)
            updates['creators']['SIM-000002'].update(incoming_summary='询问退样运费。',message_basis=c.basis(self.state,'SIM-000002'))
            c.adopt_document(self.state,source,self.book,updates)
        history=w.locate(self.book,'SIM-000002')[1]['往来邮件梗概']
        self.assertEqual(history.count('2026-10-22｜对方：询问退样运费。'),1)

    def test_both_tabs_are_checked_for_existing_contacts_and_duplicate_ids(self):
        profiles,emails=pool.known_contacts(self.state,self.book)
        self.assertIn('sim2@example.invalid',emails);self.assertIn(c.daily.profile('https://www.tiktok.com/@sim_2'),profiles)
        with self.assertRaises(c.mail.DraftError):
            w.update_fields(self.book,'SIM-000001',{'达人ID':'SIM-000002'},expected_hash=c.mail.digest(self.book.read_bytes()))
        self.assertEqual(len(w.read(self.book)['creators']),2)
        self.assertIn(w.locate(self.book,'SIM-000001')[1]['综合评分'],(None,''))

    def test_new_rows_use_main_current_headers_score_formula_and_extend_filter(self):
        candidate=support.candidate(3);candidate['达人ID']='SIM-000003'
        value={'date':'2026-10-22','batch':'sim-new','reserved':[candidate],'creators':[]}
        c.daily.recover_reserved(self.state,value,self.book)
        sheet,row=w.locate(self.book,'SIM-000003')
        self.assertEqual(sheet['name'],w.MASTER);self.assertEqual(row['综合评分'],4.5)
        self.assertEqual(row['当前建联状况'],'待首联确认');self.assertFalse(row.get('初次建联日期'))
        self.assertEqual(sheet['headers'],w.COMPACT_HEADERS)
        with ZipFile(self.book) as z:
            root=ET.fromstring(z.read(w.main_part(z)))
            self.assertEqual(root.find('s:autoFilter',w.NS).get('ref'),'A1:AB3')
        self.assertEqual(w.locate(self.book,'SIM-000002')[0]['name'],w.SHEETS[1])

    def test_mentor_document_has_all_four_sections_and_exact_incoming(self):
        source=self.report();c.adopt_document(self.state,source,self.book,self.updates(source))
        snapshot=c.model(self.state,'2026-10-22',self.book)
        from scripts.review_docx import write_document
        target=self.root/'result.docx';write_document(snapshot,target)
        text='\n'.join(p.text for p in Document(target).paragraphs)
        for label in ('对方来信原文','来信要点','我方回复要点及应对策略','我方完整回复正文'):
            self.assertEqual(text.count(label),1)
        self.assertIn(self.incoming['raw_text'],text);self.assertIn('We suggest a one-week trial.',text)

    def test_unresolved_comment_preserves_human_document_without_blocking_review_page(self):
        source=self.report();c.adopt_document(self.state,source,self.book,self.updates(source))
        report=c.build(self.state,'2026-10-22',self.book)
        document=Path(report['document']);doc=Document(document)
        p=next(p for p in doc.paragraphs if p.text=='We suggest a one-week trial.')
        doc.add_comment(p.runs[0],text='请先核实新的试用时长。',author='Mentor');doc.save(document)
        before=document.read_bytes()
        refreshed=c.build(self.state,'2026-10-22',self.book)
        self.assertTrue(refreshed['document_has_unresolved_comments'])
        self.assertEqual(document.read_bytes(),before)
        self.assertTrue(Path(refreshed['page']).exists())

if __name__=='__main__':unittest.main()
