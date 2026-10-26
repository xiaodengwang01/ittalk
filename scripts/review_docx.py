"""Mentor memo generated from the same snapshot as the private review page."""
from __future__ import annotations

from pathlib import Path
import re


def heading(text):
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s\u4e00-\u9fff]", " ", str(text))).strip()


def write_document(snapshot, output):
    from docx import Document
    from docx.shared import Cm, Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.65)
    section.left_margin = section.right_margin = Cm(1.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for key in ("ascii", "hAnsi", "cs"): fonts.set(qn("w:" + key), "Arial")
        fonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"): fonts.attrib.pop(qn("w:" + key), None)
        for border in style.element.findall(".//" + qn("w:pBdr")): border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.12
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(11.5)
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].paragraph_format.keep_with_next = True
    brief = snapshot.get("mentor_brief") or {}
    title = brief.get("title", "KOL 跟进与待决事项")
    doc.core_properties.title = title
    doc.core_properties.author = "Tailki"
    doc.add_paragraph(title, "Title")
    paragraph = doc.add_paragraph(snapshot["date"] + "  带教审核")
    paragraph.paragraph_format.space_after = Pt(4)
    doc.add_paragraph(brief.get("intro", "带教，我整理了目前待决定的合作条件和已准备的回复稿。请您确认推进方向，邮件随后仍由我核对最终版本再发送。以下沿用已采集的沟通资料；新来信和条件变化会另行更新。"))
    if brief:
        doc.add_heading("请带教决策", 1)
        for item in brief.get("decisions", []):
            doc.add_heading(item["title"], 2)
            doc.add_paragraph(item["context"])
            if item.get("recommendation"): doc.add_paragraph("我的建议：" + item["recommendation"])
            for question in item.get("questions", []):
                doc.add_paragraph("请带教确认：" + question)
        if brief.get("closure_recommendations"):
            paragraph = doc.add_heading("建议终止的合作", 1)
            paragraph.paragraph_format.page_break_before = True
            for item in brief["closure_recommendations"]:
                doc.add_heading(item["title"], 2)
                doc.add_paragraph(item["context"])
                doc.add_paragraph("建议：" + item["recommendation"])
                if item.get("handling"): doc.add_paragraph("目前处理：" + item["handling"])
        if brief.get("waiting"):
            paragraph = doc.add_heading("已回复，等待对方反馈", 1)
            if not brief.get("closure_recommendations"):
                paragraph.paragraph_format.page_break_before = True
            for item in brief["waiting"]:
                doc.add_heading(item["title"], 2)
                doc.add_paragraph(item["context"])
                if item.get("next_step"): doc.add_paragraph("下一步：" + item["next_step"])
        if brief.get("closing_note"): doc.add_paragraph(brief["closing_note"])
        # Keep the decision overview compact without reducing the body font.
        for paragraph in doc.paragraphs:
            if paragraph.style.name == "Normal":
                paragraph.paragraph_format.space_after = Pt(3)
            elif paragraph.style.name == "Heading 2":
                paragraph.paragraph_format.space_before = Pt(8)
                paragraph.paragraph_format.space_after = Pt(4)
    if snapshot["added"]:
        doc.add_heading("今日新增达人", 1)
        table = doc.add_table(rows=1, cols=3)
        table.autofit = False
        for cell, text, width in zip(table.rows[0].cells, ("达人", "合作邮箱", "内容分类"), (5, 7.4, 5)):
            cell.text = text; cell.width = Cm(width)
        for row in snapshot["added"]:
            for cell, value in zip(table.add_row().cells, (row.get("达人昵称", ""), row.get("Email", ""), row.get("首联内容分类", ""))):
                cell.text = str(value)
        for index, row in enumerate(table.rows):
            for cell in row.cells:
                properties = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    line = OxmlElement("w:" + edge)
                    for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")): line.set(qn("w:" + key), value)
                    borders.append(line)
                properties.append(borders)
                if index == 0:
                    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E9EDF2"); properties.append(shade)
                    for run in cell.paragraphs[0].runs: run.bold = True
        doc.add_paragraph("首联邮件的全部收件人、主题和完整正文见今日审核入口。")
    cards = snapshot["cards"]
    for index, card in enumerate(cards):
        paragraph = doc.add_heading(heading(card["title"]), 1)
        if index or snapshot["added"] or brief:
            paragraph.paragraph_format.page_break_before = True
        doc.add_paragraph("达人ID：" + card["creator_id"])
        doc.add_heading("对方来信原文", 2)
        doc.add_paragraph(card.get("raw_text") or "本轮来信原文尚未提供。")
        doc.add_heading("来信要点", 2)
        doc.add_paragraph(card.get("incoming_points") or card["summary"] or "待整理。")
        doc.add_heading("我方回复要点及应对策略", 2)
        doc.add_paragraph(card.get("reply_points") or card["strategy"] or "待结合来信整理。")
        for decision in card["decisions"]: doc.add_paragraph("待确认：" + str(decision))
        for part in card["details"]:
            doc.add_paragraph(part["text"])
        doc.add_heading("我方完整回复正文", 2)
        if card["draft"]:
            draft = card["draft"]
            doc.add_paragraph("收件人：" + draft["recipient"])
            for key, label in (("cc", "抄送："), ("bcc", "密送：")):
                if draft.get(key): doc.add_paragraph(label + draft[key])
            doc.add_paragraph("主题：" + draft["subject"])
            for block in draft["blocks"]:
                if "text" in block:
                    doc.add_paragraph(block["text"])
                else:
                    paragraph = doc.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_after = Pt(2)
                    images = block.get("images", [])
                    for image in images:
                        width = 2.65 if len(images) == 2 and image == images[0] else 3.25 if len(images) == 2 else 5.7 / max(len(images), 1)
                        shape = paragraph.add_run().add_picture(image["path"], width=Inches(width))
                        shape._inline.docPr.set("descr", "Tailki 产品或佩戴示意")
            paragraph = doc.add_paragraph("审核状态：待确认，未发送。")
            for run in paragraph.runs: run.font.size = Pt(9)
        else:
            doc.add_paragraph("合作邮箱：" + card["email"])
            doc.add_paragraph("本轮完整回复尚未整理。")
            if card["stale"]: doc.add_paragraph("此前汇报或回复稿的依据已有变化，需先核实最新内容。")
    if not cards and not snapshot["added"]: doc.add_paragraph("目前没有已整理的新增达人或带教待决事项。")
    if snapshot["warnings"] and brief.get("include_operational_warnings", True):
        doc.add_heading("待核实事项", 1)
        for warning in snapshot["warnings"]: doc.add_paragraph(warning)
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
